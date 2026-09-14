"""Excel control sheet: read queue, write status, green-highlight 100%-verified rows."""
import os
from datetime import datetime, timezone
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter

from . import config

HEADERS = [
    "URL", "Title override (optional)", "Short timestamp (optional)",
    "Status", "Long video link", "Short link",
    "Verified at (UTC)", "Attempts", "Error / notes",
]

COL = {name: i + 1 for i, name in enumerate(HEADERS)}

GREEN_FILL  = PatternFill("solid", start_color="C6EFCE")
GREEN_FONT  = Font(color="006100", bold=True)
RED_FILL    = PatternFill("solid", start_color="FFC7CE")
RED_FONT    = Font(color="9C0006", bold=True)
AMBER_FONT  = Font(color="9C6500", bold=True)


class Sheet:
    def __init__(self, path: str = None):
        self.path = path or config.SHEET_PATH
        if not os.path.exists(self.path):
            raise FileNotFoundError(f"Control sheet not found: {self.path}")
        self.wb = load_workbook(self.path)
        self.ws = self.wb.active
        self._validate_headers()

    def _validate_headers(self):
        row1 = [c.value for c in self.ws[1]]
        missing = [h for h in HEADERS if h not in row1]
        if missing:
            raise ValueError(f"videos.xlsx is missing columns: {missing}. Expected first row: {HEADERS}")

    def _url_col(self) -> int:
        return next(i for i, c in enumerate(self.ws[1], start=1) if c.value == "URL")

    def rows(self):
        """Yield (row_number, values_dict) for every non-empty row below the header."""
        url_col = self._url_col()
        for r in range(2, self.ws.max_row + 1):
            url = self.ws.cell(row=r, column=url_col).value
            if url and str(url).strip().startswith("http"):
                yield r, {
                    "url": str(url).strip(),
                    "title_override": self.ws.cell(row=r, column=self._c("Title override (optional)")).value,
                    "short_ts": self.ws.cell(row=r, column=self._c("Short timestamp (optional)")).value,
                    "status": self.ws.cell(row=r, column=self._c("Status")).value,
                    "long_link": self.ws.cell(row=r, column=self._c("Long video link")).value,
                    "short_link": self.ws.cell(row=r, column=self._c("Short link")).value,
                    "attempts": self.ws.cell(row=r, column=self._c("Attempts")).value or 0,
                }

    def _c(self, header: str) -> int:
        return next(i for i, c in enumerate(self.ws[1], start=1) if c.value == header)

    def set(self, row: int, header: str, value):
        self.ws.cell(row=row, column=self._c(header), value=value)

    # ---- status helpers ----
    # statuses a crashed/killed run may leave behind - must stay resumable
    LONG_RESUMABLE = {"", "pending", "⏬ downloading", "⏬ downloading"}
    SHORT_RESUMABLE = {
        "short_pending", "🟡 long uploaded - short queued",
        "✂️ building short", "🔎 verifying uploads",
        "🧪 dry-run: long skipped upload", "🧪 dry-run: short rendered, upload skipped",
    }

    def next_pending(self):
        """First row ready for the long slot (fresh or stalled mid-download)."""
        for r, v in self.rows():
            s = str(v["status"] or "").strip().lower()
            if s in {x.lower() for x in self.LONG_RESUMABLE}:
                return r, v
        return None, None

    def next_short_pending(self):
        for r, v in self.rows():
            s = str(v["status"] or "").strip().lower()
            if s in {x.lower() for x in self.SHORT_RESUMABLE}:
                return r, v
        return None, None

    def bump_attempts(self, row: int) -> int:
        n = int(self.ws.cell(row=row, column=self._c("Attempts")).value or 0) + 1
        self.set(row, "Attempts", n)
        return n

    def set_status(self, row: int, text: str):
        self.set(row, "Status", text)
        cell = self.ws.cell(row=row, column=self._c("Status"))
        cell.fill = PatternFill(fill_type=None)
        cell.font = AMBER_FONT

    def mark_verified(self, row: int, long_link: str, short_link: str):
        """100% confirmed (YouTube reports 'processed') -> green highlight the whole row."""
        now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
        self.set(row, "Status", "✅ verified")
        self.set(row, "Long video link", long_link)
        if short_link:
            self.set(row, "Short link", short_link)
        self.set(row, "Verified at (UTC)", now)
        self.set(row, "Error / notes", "")
        for col in range(1, len(HEADERS) + 1):
            cell = self.ws.cell(row=row, column=col)
            cell.fill = GREEN_FILL
            cell.font = GREEN_FONT

    def mark_failed(self, row: int, error: str):
        self.set(row, "Status", "❌ failed")
        self.set(row, "Error / notes", str(error)[:500])
        for col in range(1, len(HEADERS) + 1):
            cell = self.ws.cell(row=row, column=col)
            cell.fill = RED_FILL
        self.ws.cell(row=row, column=self._c("Status")).font = RED_FONT

    def save(self):
        self.wb.save(self.path)
